# -*- coding: utf-8 -*-
import os
import sys
import io
import json
import base64
import cv2
import numpy as np
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
import openpyxl
from datetime import datetime
from PIL import Image
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

if sys.platform.startswith('win'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except AttributeError:
        pass

# ==============================================================================
# 系統與頁面初始化
# ==============================================================================
st.set_page_config(
    page_title="SLOPEGUARD PRO // 邊坡智慧分級",
    page_icon="⛰",
    layout="wide",
    initial_sidebar_state="expanded"
)

logo_b64 = ""
for potential_path in ["LOGO.webp", "logo.webp", "LOGO.png", "logo.png"]:
    if os.path.exists(potential_path):
        try:
            with open(potential_path, "rb") as f:
                logo_b64 = base64.b64encode(f.read()).decode('utf-8')
            break
        except Exception:
            pass
logo_img_html = f'<img src="data:image/webp;base64,{logo_b64}" style="height:18px; vertical-align:middle; margin-right:6px;" />' if logo_b64 else ""

# ==============================================================================
# CSS 響應式佈局
# ==============================================================================
st.markdown("""
<style>
    .block-container { padding-top: 1.2rem !important; padding-bottom: 3.5rem !important; }
    header { visibility: hidden; }
    footer { visibility: hidden; }
    
    .app-title { font-size: 22px; font-weight: 900; margin-bottom: 4px; color: var(--text-color); line-height: 1.2; }
    .app-subtitle { font-size: 12px; color: gray; margin-bottom: 12px; letter-spacing: 0.5px; }
    
    .metrics-wrapper {
        display: grid;
        grid-template-columns: repeat(2, 1fr);
        gap: 8px;
        margin-bottom: 16px;
    }
    @media (min-width: 768px) {
        .app-title { font-size: 28px; }
        .app-subtitle { font-size: 14px; margin-bottom: 20px; }
        .metrics-wrapper { grid-template-columns: repeat(4, 1fr); gap: 12px; }
    }
    
    .metric-card {
        background-color: var(--secondary-background-color); 
        border: 1px solid var(--border-color);
        padding: 8px 10px; border-radius: 6px; border-left: 4px solid #00E5FF;
        display: flex; flex-direction: column; justify-content: center;
    }
    .metric-title { color: var(--text-color); font-size: 11px; font-weight: bold; margin-bottom: 2px; opacity: 0.8; }
    .metric-val { color: var(--text-color); font-size: 16px; font-weight: 800; font-family: monospace; line-height: 1.2; }
    
    .fixed-footer {
        position: fixed; left: 0; bottom: 0; width: 100%;
        background-color: var(--secondary-background-color); color: var(--text-color); text-align: center;
        padding: 6px 0; font-size: 11px; font-weight: bold;
        border-top: 1px solid var(--border-color); z-index: 99999;
        display: flex; justify-content: center; align-items: center;
    }
</style>
""", unsafe_allow_html=True)

# ==============================================================================
# 邊坡分析核心引擎
# ==============================================================================
class SlopeVisionEngine:
    @staticmethod
    def hex_to_bgr(hex_str):
        hex_str = hex_str.lstrip('#')
        return (int(hex_str[4:6], 16), int(hex_str[2:4], 16), int(hex_str[0:2], 16))

    @staticmethod
    def draw_styled_polygon(img, pts, color_bgr=(0, 0, 255), thickness=4, line_style="dashed"):
        num_pts = len(pts)
        if line_style == "solid":
            cv2.polylines(img, [pts], isClosed=True, color=color_bgr, thickness=thickness, lineType=cv2.LINE_AA)
            return

        pattern = [28, 8, 6, 8] if line_style == "centerline" else [5, 6] if line_style == "dotted" else [20, 12]

        for i in range(num_pts):
            p1 = pts[i]
            p2 = pts[(i + 1) % num_pts]
            dx, dy = float(p2[0] - p1[0]), float(p2[1] - p1[1])
            dist = (dx**2 + dy**2)**0.5
            if dist == 0: continue
            vx, vy = dx / dist, dy / dist

            curr, pat_idx = 0.0, 0
            while curr < dist:
                seg_len = pattern[pat_idx % len(pattern)]
                is_draw = (pat_idx % 2 == 0)
                next_d = min(curr + seg_len, dist)

                if is_draw:
                    sp = (int(p1[0] + vx * curr), int(p1[1] + vy * curr))
                    ep = (int(p1[0] + vx * next_d), int(p1[1] + vy * next_d))
                    cv2.line(img, sp, ep, color_bgr, thickness, cv2.LINE_AA)

                curr = next_d
                pat_idx += 1

    @classmethod
    def analyze(cls, cv_img, sensitivity, min_area, suppress_road, road_dilate, y_range, x_range):
        h, w = cv_img.shape[:2]
        total_pixels = h * w
        gray = cv2.cvtColor(cv_img, cv2.COLOR_BGR2GRAY)
        hsv = cv2.cvtColor(cv_img, cv2.COLOR_BGR2HSV)
        h_chan, s_chan, v_chan = cv2.split(hsv)

        b, g, r = cv_img[:, :, 0].astype(np.float32), cv_img[:, :, 1].astype(np.float32), cv_img[:, :, 2].astype(np.float32)
        nexg = (2.0 * g - r - b) / (2.0 * g + r + b + 1e-5)
        
        forest_mask = np.where((nexg > 0.015) | ((h_chan >= 20) & (h_chan <= 95) & (s_chan >= 20)), 255, 0).astype(np.uint8)
        k_forest = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (15, 15))
        forest_mask = cv2.morphologyEx(forest_mask, cv2.MORPH_CLOSE, k_forest)
        veg_coverage = (cv2.countNonZero(forest_mask) / total_pixels) * 100.0

        mask_soil = cv2.inRange(hsv, np.array([3, 15, 30]), np.array([32, 255, 250]))
        mask_structure = cv2.inRange(v_chan, int(max(70, 110 - sensitivity * 0.3)), 255)
        color_candidates = cv2.bitwise_or(mask_soil, mask_structure)

        edges = cv2.Canny(gray, 40, 120)
        kernel_edge = cv2.getStructuringElement(cv2.MORPH_RECT, (3, 3))
        edge_guided = cv2.dilate(edges, kernel_edge, iterations=1)

        hazard_raw = cv2.bitwise_and(color_candidates, edge_guided)
        hazard_filtered = cv2.bitwise_and(hazard_raw, cv2.bitwise_not(forest_mask))

        roi_filter = np.zeros((h, w), dtype=np.uint8)
        roi_filter[int(h * (y_range[0] / 100.0)):int(h * (y_range[1] / 100.0)), 
                   int(w * (x_range[0] / 100.0)):int(w * (x_range[1] / 100.0))] = 255
        hazard_filtered = cv2.bitwise_and(hazard_filtered, roi_filter)

        bridge_w, bridge_h = int(max(25, w * 0.012)), int(max(35, h * 0.02))
        kernel_bridge = cv2.getStructuringElement(cv2.MORPH_RECT, (bridge_w, bridge_h))
        scar_closed = cv2.morphologyEx(hazard_filtered, cv2.MORPH_CLOSE, kernel_bridge)
        
        kernel_open = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (11, 11))
        scar_clean = cv2.morphologyEx(scar_closed, cv2.MORPH_OPEN, kernel_open)

        contours, _ = cv2.findContours(scar_clean, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        suspicious_polys, total_scar_area = [], 0
        min_pixels = total_pixels * (min_area / 100.0)

        for cnt in contours:
            area = cv2.contourArea(cnt)
            if area > min_pixels:
                bx, by, bw, bh = cv2.boundingRect(cnt)
                if bw > (w * 0.8) and bh < (h * 0.04): continue
                epsilon = 0.005 * cv2.arcLength(cnt, True)
                poly = cv2.approxPolyDP(cnt, epsilon, True)
                suspicious_polys.append(poly.reshape(-1, 2))
                total_scar_area += area

        scar_ratio = (total_scar_area / total_pixels) * 100.0
        if scar_ratio > 4.0: status, color = "CRITICAL (崩滑)", "#FF5252"
        elif scar_ratio > 0.8: status, color = "WARNING (局部)", "#FFB300"
        else: status, color = "NORMAL (穩定)", "#00E676"

        metrics = {
            "w": w, "h": h, "res": f"{w} x {h}",
            "veg": f"{veg_coverage:.2f} %", "scar": f"{scar_ratio:.2f} %",
            "count": f"{len(suspicious_polys)} 處", "status": status, "color": color
        }
        return suspicious_polys, metrics

    @classmethod
    def render_polygons(cls, base_img, polys, color_hex="#FF0000", line_style="dashed", thickness=4):
        out = base_img.copy()
        color_bgr = cls.hex_to_bgr(color_hex)
        for pts in polys:
            cls.draw_styled_polygon(out, pts, color_bgr=color_bgr, thickness=thickness, line_style=line_style)
        return out


# ==============================================================================
# 智慧資料反轉引擎 (處理雙欄公文轉回一維資料)
# ==============================================================================
def load_and_clean_db(file_obj_or_path):
    df_preview = pd.read_excel(file_obj_or_path)
    
    if len(df_preview.columns) >= 11 or "蘇花公路邊坡分級管制樁號位置表" in df_preview.columns:
        df_raw = pd.read_excel(file_obj_or_path, header=None)
        
        header_row_idx = -1
        for i in range(min(10, len(df_raw))):
            row_vals = [str(x).strip() for x in df_raw.iloc[i].values]
            if '里程樁' in row_vals and '分級' in row_vals:
                header_row_idx = i
                break
                
        if header_row_idx != -1:
            data_rows = df_raw.iloc[header_row_idx+1:].copy()
            left_df = data_rows.iloc[:, 0:6].copy()
            left_df.columns = ["編號", "里程樁", "順/逆樁", "分級", "發生日期", "預計調降日期"]
            right_df = data_rows.iloc[:, 6:12].copy()
            right_df.columns = ["編號", "里程樁", "順/逆樁", "分級", "發生日期", "預計調降日期"]
            
            flat_df = pd.concat([left_df, right_df], ignore_index=True)
            flat_df.dropna(subset=['編號', '里程樁'], how='all', inplace=True)
            flat_df['編號'] = pd.to_numeric(flat_df['編號'], errors='coerce')
            flat_df.dropna(subset=['編號'], inplace=True)
            flat_df.sort_values(by='編號', inplace=True)
            return flat_df.reset_index(drop=True)
            
    return df_preview

# ==============================================================================
# 報表匯出引擎 
# ==============================================================================
class SlopeReportGenerator:
    @staticmethod
    def build_excel(img_before, img_after, metrics, polys, w_orig, h_orig,
                    route_name, start_mil, end_mil, spec_mil, slope_dir, orig_level, new_level,
                    cost_val, overview_desc, suggest_desc, memo_desc):
        output_excel = io.BytesIO()
        with pd.ExcelWriter(output_excel, engine='openpyxl') as writer:
            export_rows = []
            for z_i, p_pts in enumerate(polys, 1):
                for v_i, pt in enumerate(p_pts, 1):
                    export_rows.append({"單元": f"Zone {z_i}", "序號": v_i, "X": int(pt[0]), "Y": int(pt[1])})
            pd.DataFrame(export_rows).to_excel(writer, index=False, sheet_name='節點坐標')

            wb = writer.book
            ws = wb.create_sheet(title='邊坡滾動分級表', index=0)

            font_title = Font(name="標楷體", size=16, bold=True)
            font_sub = Font(name="標楷體", size=14, bold=True)
            font_header = Font(name="標楷體", size=11, bold=True)
            font_content = Font(name="標楷體", size=11)
            font_note = Font(name="標楷體", size=10, italic=True)

            fill_header = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")
            thin = Side(border_style="thin", color="000000")
            border_cell = Border(left=thin, right=thin, top=thin, bottom=thin)

            col_widths = {'A': 15, 'B': 17, 'C': 17, 'D': 17, 'E': 15, 'F': 17, 'G': 17, 'H': 17}
            for col, w in col_widths.items(): ws.column_dimensions[col].width = w

            ws.merge_cells('A1:H1')
            ws['A1'] = "交通部公路局東區養護工程分局"
            ws['A1'].font = font_title
            
            ws.merge_cells('A2:H2')
            ws['A2'] = "邊坡滾動分級表"
            ws['A2'].font = font_sub

            ws['A3'] = "養護單位"
            ws.merge_cells('B3:D3')
            ws['B3'] = "東區養護工程分局南澳工務段"
            ws['E3'] = "邊坡\n原有分級"
            ws.merge_cells('F3:H3')
            ws['F3'] = orig_level

            ws['A4'] = "邊坡位置"
            ws.merge_cells('B4:D4')
            if spec_mil: ws['B4'] = f"{route_name}{start_mil}~{end_mil}\n({route_name}{spec_mil})"
            else: ws['B4'] = f"{route_name}{start_mil}~{end_mil}"
            ws['E4'] = "里 程"
            ws.merge_cells('F4:H4')
            box_shun = "■" if "順" in slope_dir else "□"
            box_ni = "■" if "逆" in slope_dir else "□"
            ws['F4'] = f"{box_shun}順樁        {box_ni}逆樁"

            ws['A5'] = "災害所需經費"
            ws.merge_cells('B5:D5')
            ws['B5'] = cost_val
            ws['E5'] = "邊坡等級\n調 整"
            ws.merge_cells('F5:H5')
            ws['F5'] = f"{orig_level}    →    {new_level}"

            ws['A6'] = "現況概述"
            ws.merge_cells('B6:H6')
            ws['B6'] = overview_desc

            ws['A7'] = "擬改善情形"
            ws.merge_cells('B7:H7')
            ws['B7'] = suggest_desc

            ws.merge_cells('A8:D8')
            ws['A8'] = "邊坡滾動前"
            ws.merge_cells('E8:H8')
            ws['E8'] = "邊坡滾動後"

            ws['A23'] = "備 註"
            ws.merge_cells('B23:H23')
            ws['B23'] = memo_desc

            ws.merge_cells('A24:D24')
            ws['A24'] = "工務段：\n\n\n\n\n聘用高級工程師 張仁瀚"
            ws.merge_cells('E24:H24')
            ws['E24'] = "主管單位：\n\n\n\n\n段長 劉錦龍"

            ws.merge_cells('A25:H25')
            ws['A25'] = "註：本報告表係用於邊坡調整適用之，1式3份報分局核定。"
            ws['A25'].font = font_note
            ws['A25'].alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[1].height = 25
            ws.row_dimensions[2].height = 25
            ws.row_dimensions[3].height = 35
            ws.row_dimensions[4].height = 45
            ws.row_dimensions[5].height = 35
            ws.row_dimensions[6].height = 60
            ws.row_dimensions[7].height = 45
            ws.row_dimensions[8].height = 25
            
            ws.merge_cells('A9:D22')
            ws.merge_cells('E9:H22')
            for r in range(9, 23): ws.row_dimensions[r].height = 23
                
            ws.row_dimensions[23].height = 45
            ws.row_dimensions[24].height = 100
            ws.row_dimensions[25].height = 20

            for r in range(1, 25):
                for col_idx in range(1, 9):
                    cell = ws.cell(row=r, column=col_idx)
                    if r >= 3: cell.border = border_cell
                    if r in [1, 2]: cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif r == 24:
                        cell.font = font_content
                        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    elif col_idx in [1, 5] and r in [3, 4, 5, 6, 7, 23]:
                        cell.fill = fill_header
                        cell.font = font_header
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    elif r == 8:
                        cell.fill = fill_header
                        cell.font = font_header
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif r >= 3:
                        cell.font = font_content
                        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            try:
                from openpyxl.drawing.image import Image as OpenpyxlImage
                target_w, target_h = 420, 315 
                if img_before is not None:
                    buf_b = io.BytesIO(cv2.imencode('.png', cv2.resize(img_before, (target_w, target_h), interpolation=cv2.INTER_AREA))[1].tobytes())
                    oxl_img_b = OpenpyxlImage(buf_b)
                    oxl_img_b.width, oxl_img_b.height = target_w, target_h
                    ws.add_image(oxl_img_b, "A9")
                if img_after is not None:
                    buf_a = io.BytesIO(cv2.imencode('.png', cv2.resize(img_after, (target_w, target_h), interpolation=cv2.INTER_AREA))[1].tobytes())
                    oxl_img_a = OpenpyxlImage(buf_a)
                    oxl_img_a.width, oxl_img_a.height = target_w, target_h
                    ws.add_image(oxl_img_a, "E9")
            except Exception as e: print(e)

        return output_excel.getvalue()

    @staticmethod
    def build_cad_dxf(polys, w_orig, h_orig, route_name, stake_mileage, line_style="dashed"):
        lines = ["0", "SECTION", "2", "HEADER", "9", "$ACADVER", "1", "AC1009", "0", "ENDSEC", "0", "SECTION", "2", "ENTITIES"]
        for z_i, poly in enumerate(polys, 1):
            lines.extend(["0", "POLYLINE", "8", "SLOPE", "66", "1", "70", "1"])
            for pt in poly: lines.extend(["0", "VERTEX", "8", "SLOPE", "10", f"{float(pt[0]):.3f}", "20", f"{float(h_orig - pt[1]):.3f}", "30", "0.0"])
            lines.extend(["0", "SEQEND"])
        lines.extend(["0", "ENDSEC", "0", "EOF"])
        return "\n".join(lines).encode('utf-8')

    @staticmethod
    def build_master_control_table_via_template(df_db, active_mileages, template_path="蘇花公路邊坡分級管制樁號位置表（範本）xlsx.xlsx"):
        if not os.path.exists(template_path):
            st.error(f"找不到範本檔：{template_path}，請確保該檔案在目錄下。")
            return None

        wb = openpyxl.load_workbook(template_path)
        ws = wb.active

        merged_ranges = list(ws.merged_cells.ranges)
        for m_range in merged_ranges:
            if m_range.min_row >= 5:
                ws.unmerge_cells(str(m_range))

        df_clean = df_db.copy()
        
        lvl_col = "定性分級" if "定性分級" in df_clean.columns else "分級"
        def extract_grade(g):
            g_str = str(g).upper()
            if 'A' in g_str: return 'A'
            if 'B' in g_str: return 'B'
            if 'C' in g_str: return 'C'
            if 'D' in g_str: return 'D'
            return 'Z'
        df_clean['SortGrade'] = df_clean[lvl_col].apply(extract_grade) if lvl_col in df_clean.columns else 'Z'
        
        def parse_mileage(m_str):
            try:
                parts = str(m_str).lower().split('k')
                if len(parts) == 2:
                    km = float(parts[0])
                    m = float(parts[1].replace('+', ''))
                    return km * 1000 + m
                return 999999.0
            except:
                return 999999.0

        if '里程樁號(起)' in df_clean.columns:
            df_clean['SortMileage'] = df_clean['里程樁號(起)'].apply(parse_mileage)
        elif '里程樁' in df_clean.columns:
            df_clean['SortMileage'] = df_clean['里程樁'].apply(parse_mileage)
        else:
            df_clean['SortMileage'] = 999999.0
            
        sort_cols = ['SortGrade']
        if '路線' in df_clean.columns: sort_cols.append('路線')
        sort_cols.append('SortMileage')
            
        df_sorted = df_clean.sort_values(by=sort_cols).reset_index(drop=True)
        
        count_A = sum(df_sorted['SortGrade'] == 'A')
        count_B = sum(df_sorted['SortGrade'] == 'B')
        count_C = sum(df_sorted['SortGrade'] == 'C')
        count_D = sum(df_sorted['SortGrade'] == 'D')
        total_count = len(df_sorted)

        ws['D3'], ws['E3'], ws['F3'], ws['G3'] = count_A, count_B, count_C, count_D
        ws['H3'] = total_count

        font_content = Font(name="標楷體", size=10)
        fill_active = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        thin = Side(border_style="thin", color="000000")
        border_cell = Border(left=thin, right=thin, top=thin, bottom=thin)

        for r in range(5, 200):
            for c in range(1, 13):
                cell = ws.cell(row=r, column=c)
                if type(cell).__name__ != 'MergedCell':
                    try:
                        cell.value = ""
                        cell.fill = PatternFill()
                    except:
                        pass

        half_len = (total_count + 1) // 2
        
        for idx, row in df_sorted.iterrows():
            is_right_col = idx >= half_len
            c_offset = 6 if is_right_col else 0
            r_idx = 5 + (idx - half_len) if is_right_col else 5 + idx

            ws.row_dimensions[r_idx].height = 16.5
            ws.cell(row=r_idx, column=c_offset+1, value=idx+1)
            
            if "里程樁" in row and pd.notna(row["里程樁"]):
                raw_full_mil = str(row["里程樁"]).replace("*", "")
            else:
                route = str(row.get("路線", ""))
                start = str(row.get("里程樁號(起)", ""))
                end = str(row.get("里程樁號(迄)", ""))
                spec = str(row.get("特定樁號", ""))
                
                raw_full_mil = f"{route}{start}~{end}"
                if spec and spec != "nan":
                    raw_full_mil += f"({spec})"
            
            is_active = False
            for act_m in active_mileages:
                clean_act = act_m.replace("*", "")
                if clean_act and (clean_act in raw_full_mil or raw_full_mil in clean_act):
                    is_active = True
                    break
            
            display_mil = f"*{raw_full_mil}" if is_active else raw_full_mil
            ws.cell(row=r_idx, column=c_offset+2, value=display_mil)

            dir_val = str(row.get("順/逆樁", row.get("方向", "")))
            if dir_val == "nan": dir_val = ""
            ws.cell(row=r_idx, column=c_offset+3, value=dir_val)

            pure_grade = row['SortGrade']
            ws.cell(row=r_idx, column=c_offset+4, value=pure_grade if pure_grade != 'Z' else "")

            date_val = str(row.get("發生日期", ""))
            if date_val == "nan": date_val = ""
            ws.cell(row=r_idx, column=c_offset+5, value=date_val)

            fix_date = str(row.get("預計調降日期", ""))
            if fix_date == "nan": fix_date = ""
            ws.cell(row=r_idx, column=c_offset+6, value=fix_date)
            
            for c in range(1, 7):
                cell = ws.cell(row=r_idx, column=c_offset+c)
                cell.border = border_cell
                cell.font = font_content
                cell.alignment = Alignment(horizontal="left" if c in [2] else "center", vertical="center")
                if is_active:
                    cell.fill = fill_active

        output_excel = io.BytesIO()
        wb.save(output_excel)
        return output_excel.getvalue()


# ==============================================================================
# 系統狀態與 Raw Data 資料庫管理
# ==============================================================================
if 'df_db' not in st.session_state:
    if os.path.exists("邊坡資料.xlsx"):
        st.session_state.df_db = load_and_clean_db("邊坡資料.xlsx")
    elif os.path.exists("蘇花公路邊坡分級管制樁號位置表（範本）xlsx.xlsx"):
        st.session_state.df_db = load_and_clean_db("蘇花公路邊坡分級管制樁號位置表（範本）xlsx.xlsx")
    else:
        st.session_state.df_db = pd.DataFrame()

if 'active_mileages' not in st.session_state:
    st.session_state.active_mileages = []

with st.sidebar:
    st.header("🗄️ Raw Data 邊坡台帳資料庫")
    
    if not st.session_state.df_db.empty:
        st.success("✅ 已載入台帳資料庫")
    
    db_file = st.file_uploader("手動更新 Raw Data 台帳覆蓋", type=["xlsx"])
    if db_file:
        try:
            st.session_state.df_db = load_and_clean_db(db_file)
            st.success("✅ 已切換並壓平為最新台帳")
        except Exception:
            st.error("讀取失敗，請確認檔案格式。")
            
    db_options = []
    if not st.session_state.df_db.empty:
        for idx, row in st.session_state.df_db.iterrows():
            if "里程樁" in row and pd.notna(row["里程樁"]):
                mil_str = str(row["里程樁"]).replace("*", "")
            else:
                r_name = str(row.get("路線", ""))
                r_start = str(row.get("里程樁號(起)", ""))
                r_end = str(row.get("里程樁號(迄)", ""))
                r_spec = str(row.get("特定樁號", ""))
                mil_str = f"{r_name}{r_start}~{r_end}"
                if str(r_spec) != "nan" and r_spec:
                    mil_str += f"({r_spec})"
            if mil_str not in ["nan", ""]: 
                db_options.append(mil_str)
            
    sel_slope = st.selectbox("選擇既有邊坡帶入公文", db_options) if db_options else None
    
    st.markdown("---")
    st.header("🎛 標註樣式與 AI 參數")
    
    color_options = {"工程紅 (#FF0000)": "#FF0000", "青藍 (#00E5FF)": "#00E5FF", "警戒黃 (#FFD600)": "#FFD600", "高對比綠 (#00E676)": "#00E676"}
    stroke_color = color_options[st.selectbox("線條顏色", list(color_options.keys()), index=0)]
    line_style_map = {"虛線": "dashed", "實線": "solid", "中心線": "centerline", "點線": "dotted"}
    active_line_style = line_style_map[st.selectbox("線條種類", list(line_style_map.keys()), index=0)]
    active_stroke_width = st.slider("粗細 (px)", 1, 12, 4)
    sensitivity = st.slider("辨識敏感度 (%)", 10, 100, 50, 5)
    min_area = st.slider("雜訊過濾閥值 (%)", 0.05, 2.0, 0.2, 0.05)


# ==============================================================================
# 主頁面 UI 
# ==============================================================================
st.markdown("""
<div class="app-title">⛰️ SLOPEGUARD PRO</div>
<div class="app-subtitle">公路邊坡分級管理系統 ｜ 邊坡滾動分級表與 CAD/Word/Excel 匯出版</div>
""", unsafe_allow_html=True)

tab_cv, tab_audit, tab_control = st.tabs(["影像辨識與滾動報表", "A級降等合規審查", "總台帳與管制位置表"])

with tab_cv:
    c_up1, c_up2 = st.columns(2)
    with c_up1: file_before = st.file_uploader("📸 上傳『邊坡滾動前』照片", type=["jpg", "jpeg", "png"])
    with c_up2: file_after = st.file_uploader("📸 上傳『邊坡滾動後』照片 (進行 AI 辨識)", type=["jpg", "jpeg", "png"])

    img_before_cv, img_after_cv, annotated_bgr = None, None, None
    metrics = {"veg": "0.00 %", "scar": "0.00 %", "color": "#00E676", "status": "NORMAL"}
    w_orig, h_orig = 1920, 1080 

    if file_before:
        img_before_cv = cv2.imdecode(np.asarray(bytearray(file_before.read()), dtype=np.uint8), cv2.IMREAD_COLOR)

    if file_after:
        img_after_cv = cv2.imdecode(np.asarray(bytearray(file_after.read()), dtype=np.uint8), cv2.IMREAD_COLOR)
        h_orig, w_orig = img_after_cv.shape[:2]

        polys_detected, metrics = SlopeVisionEngine.analyze(
            img_after_cv, sensitivity, min_area, False, 2, (0, 100), (0, 100)
        )

        if "user_polys" not in st.session_state or st.session_state.get("last_file") != file_after.name:
            st.session_state.user_polys = [p.copy() for p in polys_detected]
            st.session_state.last_file = file_after.name

        metrics_html = f"""
        <div class="metrics-wrapper">
            <div class="metric-card" style="border-left-color: #00E5FF;">
                <div class="metric-title">植生覆蓋比例</div>
                <div class="metric-val">{metrics['veg']}</div>
            </div>
            <div class="metric-card" style="border-left-color: #FFB300;">
                <div class="metric-title">裸露崩塌比例</div>
                <div class="metric-val">{metrics['scar']}</div>
            </div>
            <div class="metric-card" style="border-left-color: #FF5252;">
                <div class="metric-title">標註破壞單元</div>
                <div class="metric-val">{len(st.session_state.user_polys)} 處</div>
            </div>
            <div class="metric-card" style="border-left-color: {metrics['color']};">
                <div class="metric-title">初判危害等級</div>
                <div class="metric-val" style="color: {metrics['color']}; font-size: 14px;">{metrics['status'].split(' ')[0]}</div>
            </div>
        </div>
        """
        st.markdown(metrics_html, unsafe_allow_html=True)

        annotated_bgr = SlopeVisionEngine.render_polygons(img_after_cv, st.session_state.user_polys, color_hex=stroke_color, line_style=active_line_style, thickness=active_stroke_width)
        
        target_display_w = 1280
        aspect_ratio = float(h_orig) / float(w_orig)
        resized_for_web = cv2.resize(img_after_cv, (target_display_w, int(target_display_w * aspect_ratio)), interpolation=cv2.INTER_AREA)
        _, buf = cv2.imencode('.jpg', resized_for_web, [cv2.IMWRITE_JPEG_QUALITY, 85])
        img_b64 = base64.b64encode(buf.tobytes()).decode('utf-8')

        polys_js = [[[float(pt[0]), float(pt[1])] for pt in p] for p in st.session_state.user_polys]

        html_template_path = "canvas_editor.html"
        if os.path.exists(html_template_path):
            with open(html_template_path, "r", encoding="utf-8") as f:
                html_content = f.read().replace("__IMG_ORIG_W__", str(w_orig)).replace("__IMG_ORIG_H__", str(h_orig))
                html_content = html_content.replace("__STROKE_COLOR__", stroke_color).replace("__LINE_STYLE__", active_line_style)
                html_content = html_content.replace("__STROKE_WIDTH__", str(active_stroke_width)).replace("__IMG_B64__", img_b64).replace("__POLYS_DATA__", json.dumps(polys_js))
            components.html(html_content, height=730)

        c_in1, c_in2 = st.columns([3, 1])
        with c_in1: json_input = st.text_input("貼上圖面最新頂點坐標資料 (JSON)", placeholder="點擊上方『📋 複製最新編輯成果』後貼於此處...", label_visibility="collapsed")
        with c_in2:
            if st.button("套用並更新圖檔與台帳", type="primary", use_container_width=True):
                if json_input:
                    try:
                        st.session_state.user_polys = [np.array(p, dtype=np.int32) for p in json.loads(json_input)]
                        st.success("已成功同步編輯成果！")
                        st.rerun()
                    except Exception: st.error("JSON 格式不符！")

    st.markdown("---")

    # =====================================================================
    # 報表資料填寫區
    # =====================================================================
    st.subheader("📑 滾動分級表資料填寫與自動更新 Raw Data")
    d_route, d_start, d_end, d_spec, d_orig_lvl, d_dir = "", "", "", "", "A", "順樁"
    today_str = datetime.now().strftime("%Y/%m/%d")
    
    row_idx = None
    if not st.session_state.df_db.empty and sel_slope:
        try:
            row_idx = db_options.index(sel_slope)
            row_data = st.session_state.df_db.iloc[row_idx]
            
            if "里程樁" in row_data and pd.notna(row_data["里程樁"]):
                d_start = str(row_data["里程樁"]).replace("*", "")
                d_route = d_start.split("K")[0] if "K" in d_start else "台9丁線" 
            else:
                d_route = str(row_data.get("路線", ""))
                d_start = str(row_data.get("里程樁號(起)", ""))
                d_end = str(row_data.get("里程樁號(迄)", ""))
                d_spec = str(row_data.get("特定樁號", ""))
                
            d_dir = str(row_data.get("順/逆樁", row_data.get("方向", "順樁")))
            lvl_col = "定性分級" if "定性分級" in st.session_state.df_db.columns else "分級"
            d_orig_lvl = str(row_data.get(lvl_col, "A"))
        except: pass

    if d_spec == "nan": d_spec = ""
    if "順" in d_dir: dir_index = 0
    elif "逆" in d_dir: dir_index = 1
    else: dir_index = 0

    cr1, cr2, cr3, cr4 = st.columns(4)
    with cr1: in_route = st.text_input("路線名稱", value=d_route)
    with cr2: in_start = st.text_input("起點樁/完整里程樁", value=d_start)
    with cr3: in_end = st.text_input("迄點樁號 (若上格已完整可留空)", value=d_end)
    with cr4: in_spec = st.text_input("特定樁號 (可留空)", value=d_spec)

    cr5, cr6, cr7, cr8 = st.columns(4)
    with cr5: in_dir = st.selectbox("里程方向", ["順樁", "逆樁"], index=dir_index)
    with cr6: in_orig_level = st.text_input("原有定性分級", value=f"{d_orig_lvl.replace('級邊坡', '')}級邊坡")
    with cr7: in_new_level = st.selectbox("調整後等級", ["A級邊坡", "B級邊坡", "C+級邊坡", "C級邊坡", "D級邊坡"], index=1)
    with cr8: in_cost = st.text_input("災害所需經費", value="初估 285,000 元")
    
    cr9, cr10 = st.columns(2)
    with cr9: in_date = st.text_input("發生日期", value=today_str)
    with cr10: in_fix_date = st.text_input("預計調降日期", value=today_str)

    in_overview = st.text_area("現況概述", value="受強降雨影響，邊坡表土及岩層遇水軟化產生坡面沖蝕，導致植被及表土流失形成大面積坡面裸露。")
    in_suggest = st.text_area("擬改善情形", value="上邊坡搶修 (型框植生)")
    in_memo = st.text_area("備註", value="上邊坡搶災。")

    cb1, cb2, cb3 = st.columns(3)
    with cb1:
        export_after = annotated_bgr if annotated_bgr is not None else img_after_cv
        excel_data = SlopeReportGenerator.build_excel(
            img_before_cv, export_after, metrics, st.session_state.get("user_polys", []), 
            w_orig, h_orig, in_route, in_start, in_end, in_spec, in_dir, 
            in_orig_level, in_new_level, in_cost, in_overview, in_suggest, in_memo
        )
        st.download_button("📊 匯出單筆滾動報表 (Excel)", excel_data, f"Slope_{in_route}_{in_start[:5]}.xlsx", use_container_width=True)
        
    with cb2:
        if st.button("🔄 將滾動分級與日期寫入記憶體 Raw Data", use_container_width=True):
            if row_idx is not None and not st.session_state.df_db.empty:
                lvl_col = "定性分級" if "定性分級" in st.session_state.df_db.columns else "分級"
                st.session_state.df_db.at[row_idx, lvl_col] = in_new_level.replace("級邊坡", "")
                st.session_state.df_db.at[row_idx, "發生日期"] = in_date
                st.session_state.df_db.at[row_idx, "預計調降日期"] = in_fix_date
                
                if "里程樁" in st.session_state.df_db.columns:
                    full_mil = str(st.session_state.df_db.at[row_idx, "里程樁"]).replace("*", "")
                else:
                    full_mil = f"{in_route}{in_start}~{in_end}"
                    if in_spec: full_mil += f"({in_spec})"
                    
                if full_mil not in st.session_state.active_mileages:
                    st.session_state.active_mileages.append(full_mil)
                    
                st.success(f"✅ 已成功將 {full_mil} 從 {d_orig_lvl} 級更新至 {in_new_level}！系統已自動將其搬移分類，請至「總台帳與管制位置表」匯出最新版範本。")
            else:
                st.error("請先於左側邊欄選擇欲滾動更新的既有邊坡！")

# ==============================================================================
# 115年公路局修訂版：嚴謹的 A級邊坡降等合規審查
# ==============================================================================
with tab_audit:
    st.markdown("### 🏛️ A級邊坡降至 C+ 或 C 級合規性審查")
    st.info("規範依據：公路局115年5月修訂版「公路邊坡分級管理機制」 (105年10月31日路養管字第1050135768號函)[cite: 13, 14, 15]")
    
    st.markdown("#### 第一階段：啟動時機與條件[cite: 15]")
    check_1 = st.checkbox("☑ 年度汛期後啟動 (已於該年度汛期正式結束後提出申請)[cite: 15]")
    check_2 = st.checkbox("☑ 完成長期觀測 (已歷經至少『1個完整汛期』的穩定性觀察與復建效能檢核)[cite: 15]")
    
    st.markdown("#### 第二階段：資料準備與專家佐證[cite: 15]")
    check_3 = st.checkbox("☑ 具備致災原因探討 (已釐清原始災害機制)[cite: 15]")
    check_4 = st.checkbox("☑ 具備削減措施說明 (已包含工程與防護成效佐證)[cite: 15]")
    check_5 = st.checkbox("☑ 具備完整佐證資料 (含安定分析結果、影像判釋、現場照片及文字降等理由)[cite: 15]")
    check_6 = st.checkbox("☑ 已由所轄分局召開專家審查會議 (邀請專家、學者或技師就致災原因與削減措施進行降級檢討)[cite: 15]")
    
    st.markdown("---")
    st.markdown("#### 第三階段：降等分流判定 (C+級 或 C級)[cite: 13, 14, 15]")
    has_remote_factor = st.radio(
        "該邊坡是否具備「遠端致災因素」？(如：需持續關注之上/下邊坡、曲流/浪襲攻擊面、坡頂人為開發、明隧道上方邊坡等)[cite: 13, 14]",
        ("否 (無遠端致災因素，符合降至 C級)", "是 (具備遠端致災因素，僅能降至 C+級)")
    )

    if st.button("執行 115 年合規性檢核", type="primary", use_container_width=True):
        if all([check_1, check_2, check_3, check_4, check_5, check_6]):
            target_level = "C+ 級" if "是" in has_remote_factor else "C 級"
            st.success(f"✅ 【審核通過】：各項法定要件均合規，專家會議佐證完備！\n\n📌 判定結果：該邊坡符合降等至 **{target_level}** 邊坡之條件。")
        else:
            st.error("❌ 【審核退回】：未滿足 115 年修訂版降等要件。請確保已歷經完整汛期、備妥工程防護佐證，並完成專家審查會議。")

with tab_control:
    st.subheader("📑 邊坡分級管制總表 (套用範本格式)")
    if not st.session_state.df_db.empty:
        st.caption("系統將讀取 `邊坡資料.xlsx` (Raw Data)，自動分群排序(A->B->C->D)並依據數字大小遞增，最後套用並填入 `蘇花公路邊坡分級管制樁號位置表（範本）xlsx.xlsx` 中！")
        active_selections = st.multiselect("手動指定本次滾動的邊坡 (將自動標註 * 與橘色底色)：", db_options, default=st.session_state.active_mileages)
        
        if st.button("生成蘇花公路邊坡分級管制樁號位置表 (載入範本版型)", type="primary"):
            master_excel = SlopeReportGenerator.build_master_control_table_via_template(st.session_state.df_db, active_selections)
            if master_excel:
                st.download_button("📥 下載最新管制總表.xlsx", master_excel, "蘇花公路邊坡分級管制樁號位置表_最新更新版.xlsx", use_container_width=True)
    else:
        st.warning("請先於左側邊欄匯入 `邊坡資料.xlsx`，確保目錄下有範本檔案。")

st.markdown(f'<div class="fixed-footer">{logo_img_html}Powered by 公路局東區養護工程分局南澳工務段</div>', unsafe_allow_html=True)